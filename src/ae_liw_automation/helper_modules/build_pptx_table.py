# build_pptx_table.py

from pandas import DataFrame
from pptx.util import Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from ae_liw_automation.helper_modules.style_table_cell import style_table_cell


def build_pptx_table(slide, table_shape_element, final_result_df_combined: DataFrame,
                     base_row: DataFrame, index_name: str, old_table_name: str, top_offset):
    rows, cols = final_result_df_combined.shape

    # Define styling properties
    header_bg_color = RGBColor(90, 128, 184)   # Dark Blue
    header_text_color = RGBColor(255, 255, 255) # White
    data_text_color = RGBColor(0, 0, 0)         # Black text
    data_bg_color = RGBColor(224, 229, 240)     # Light blue for data rows

    # Step 1: Remove existing table
    slide.shapes._spTree.remove(table_shape_element)

    # Step 2: Add a new table to the slide
    table_shape_object = slide.shapes.add_table(rows + 1, cols + 1, Inches(.5), top_offset, Inches(6.5), Inches(2.5))
    table_shape_object.name = old_table_name
    table_shape = table_shape_object.table

    # Step 3: Insert column headers (including index column)
    style_table_cell(table_shape.cell(0, 0),
                     text=index_name,
                     font_size=12,
                     bold=True,
                     color=header_text_color,
                     bg_color=header_bg_color,
                     h_alignment=PP_ALIGN.CENTER,
                     v_alignment=MSO_ANCHOR.MIDDLE,
                     )
    for col_idx, col_name in enumerate(final_result_df_combined.columns):
        style_table_cell(table_shape.cell(0, col_idx + 1),
                         col_name,
                         font_size=12,
                         bold=True,
                         color=header_text_color,
                         bg_color=header_bg_color,
                         h_alignment=PP_ALIGN.CENTER,
                         v_alignment=MSO_ANCHOR.MIDDLE,
                         )

    # Step 4: Insert data rows (including index values)
    for row_idx, (index_value, row) in enumerate(final_result_df_combined.iterrows()):
        style_table_cell(table_shape.cell(row_idx + 1, 0),
                         str(index_value),
                         font_size=12,
                         bold=False,
                         color=data_text_color,
                         bg_color=data_bg_color,
                         h_alignment=PP_ALIGN.LEFT,
                         v_alignment=MSO_ANCHOR.MIDDLE,
                         )
        for col_idx, value in enumerate(row):
            style_table_cell(table_shape.cell(row_idx + 1, col_idx + 1),
                             str(value),
                             font_size=12,
                             bold=False,
                             color=data_text_color,
                             bg_color=data_bg_color,
                             h_alignment=PP_ALIGN.CENTER,
                             v_alignment=MSO_ANCHOR.MIDDLE,
                             )

    # Step 5: Style the last row (Base:) with header styling
    style_table_cell(table_shape.cell(len(final_result_df_combined), 0),
                     text=base_row.index[0],
                     font_size=12,
                     bold=True,
                     color=header_text_color,
                     bg_color=header_bg_color,
                     h_alignment=PP_ALIGN.LEFT,
                     v_alignment=MSO_ANCHOR.MIDDLE,
                     )
    for col_number, value in enumerate(final_result_df_combined.loc['Base:']):
        style_table_cell(table_shape.cell(len(final_result_df_combined), col_number + 1),
                         font_size=12,
                         bold=True,
                         color=header_text_color,
                         bg_color=header_bg_color,
                         h_alignment=PP_ALIGN.CENTER,
                         v_alignment=MSO_ANCHOR.MIDDLE,
                         )
